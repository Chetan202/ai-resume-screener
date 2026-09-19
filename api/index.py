import io
import json
import os
import time
from dotenv import load_dotenv
import re
import httpx
from pydantic import BaseModel
from bs4 import BeautifulSoup
class JobUrlRequest(BaseModel):
    url: str


class JobUrlResponse(BaseModel):
    url: str
    title: str | None
    job_description: str
load_dotenv()
from docx import Document
from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from groq import Groq
from pydantic import BaseModel, Field
from pypdf import PdfReader

MAX_RESUMES = 20
MAX_FILE_SIZE_MB = 10
REQUESTS_PER_MINUTE = 10

SKILLS_WEIGHT = 50
EXPERIENCE_WEIGHT = 30
EDUCATION_WEIGHT = 20

MODEL_NAME = os.getenv("GROQ_MODEL", "openai/gpt-oss-120b")


class JobDescription(BaseModel):
    role: str
    required_skills: list[str]
    preferred_skills: list[str]
    minimum_experience: float | None
    education_requirements: list[str]
    responsibilities: list[str]


class Experience(BaseModel):
    company: str | None = None
    role: str | None = None
    duration: str | None = None
    description: str | None = None
    skills_used: list[str] = []


class Resume(BaseModel):
    name: str | None = None
    email: str | None = None
    phone: str | None = None
    total_experience_years: float | None = None
    skills: list[str] = []
    experiences: list[Experience] = []
    education: list[str] = []
    projects: list[str] = []
    certifications: list[str] = []


class MatchResult(BaseModel):
    skills_percentage: float = Field(ge=0, le=100)
    experience_percentage: float = Field(ge=0, le=100)
    education_percentage: float = Field(ge=0, le=100)
    matched_skills: list[str]
    missing_skills: list[str]
    experience_match: bool
    education_match: bool
    reasons: list[str]


class FinalResult(BaseModel):
    file_name: str
    name: str | None
    email: str | None = None
    phone: str | None = None
    score: float = Field(ge=0, le=100)
    verdict: str
    matched_skills: list[str]
    missing_skills: list[str]
    experience_match: bool
    education_match: bool
    reasons: list[str]


class ScreenResponse(BaseModel):
    role: str
    required_skills: list[str]
    results: list[FinalResult]
    skipped: list[str]


app = FastAPI(title="Resume Screener")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)

request_history: dict[str, list[float]] = {}


def get_client():
    api_key = os.getenv("GROQ_API_KEY")

    if not api_key:
        raise HTTPException(
            status_code=500,
            detail="GROQ_API_KEY is not configured on the server.",
        )

    return Groq(api_key=api_key)


def check_rate_limit(user_id):
    current_time = time.time()
    history = request_history.get(user_id, [])

    history = [
        timestamp
        for timestamp in history
        if current_time - timestamp < 60
    ]

    if len(history) >= REQUESTS_PER_MINUTE:
        raise HTTPException(
            status_code=429,
            detail=f"Rate limit exceeded. Maximum {REQUESTS_PER_MINUTE} requests per minute.",
        )

    history.append(current_time)
    request_history[user_id] = history


def read_pdf(data):
    reader = PdfReader(io.BytesIO(data))
    text = ""

    for page in reader.pages:
        page_text = page.extract_text()

        if page_text:
            text += page_text + "\n"

    return text


def read_docx(data):
    document = Document(io.BytesIO(data))
    text = ""

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text += paragraph.text + "\n"

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text += cell.text + "\n"

    return text


def read_resume(file_name, data):
    lowered = file_name.lower()

    if lowered.endswith(".pdf"):
        return read_pdf(data)

    if lowered.endswith(".docx"):
        return read_docx(data)

    raise HTTPException(
        status_code=400,
        detail=f"{file_name}: only PDF and DOCX files are supported.",
    )


def complete_json(client, prompt, schema):
    response = client.chat.completions.create(
        model=MODEL_NAME,
        messages=[{"role": "user", "content": prompt}],
        response_format={"type": "json_object", "schema": schema},
    )

    content = response.choices[0].message.content

    return json.loads(content)


def parse_job_description(client, job_description):
    schema = JobDescription.model_json_schema()

    prompt = f"""
Extract structured information from this job description.

Return ONLY JSON matching this schema:

{schema}

Rules:
- Do not invent information.
- If minimum experience is not mentioned, return null.
- If a list has no information, return an empty list.

Job Description:

{job_description}
"""

    return JobDescription(**complete_json(client, prompt, schema))


def parse_resume(client, resume_text):
    schema = Resume.model_json_schema()

    prompt = f"""
You are an expert resume parser.

Extract structured information from the resume.

Return ONLY JSON matching this schema:

{schema}

Rules:
- Do not invent information.
- If information is unavailable, return null.
- If a list has no information, return an empty list.
- Include internships inside experiences.
- Extract skills mentioned throughout the resume.
- Calculate total experience only when it can be reasonably determined from the resume.

Resume:

{resume_text}
"""

    return Resume(**complete_json(client, prompt, schema))


def match_resume(client, job, resume):
    schema = MatchResult.model_json_schema()

    prompt = f"""
Compare the candidate against the job description.

Job Description:

{job.model_dump_json(indent=2)}

Candidate Resume:

{resume.model_dump_json(indent=2)}

Return ONLY JSON matching this schema:

{schema}

Rules:

1. Compare required skills with candidate skills.
2. Calculate skills percentage from 0 to 100.
3. Determine whether the experience requirement is satisfied.
4. Calculate experience percentage from 0 to 100.
5. Determine whether the education requirement is satisfied.
6. Calculate education percentage from 0 to 100.
7. List matched skills.
8. List missing required skills.
9. Give short reasons.
10. Do not invent candidate information.
"""

    return MatchResult(**complete_json(client, prompt, schema))


def calculate_score(match):
    score = (
        match.skills_percentage * SKILLS_WEIGHT / 100
        + match.experience_percentage * EXPERIENCE_WEIGHT / 100
        + match.education_percentage * EDUCATION_WEIGHT / 100
    )

    return round(max(0, min(100, score)), 2)


def get_verdict(score):
    if score >= 80:
        return "Strong Match"

    if score >= 60:
        return "Moderate Match"

    if score >= 40:
        return "Weak Match"

    return "Poor Match"


def create_final_result(file_name, resume, match):
    score = calculate_score(match)

    return FinalResult(
        file_name=file_name,
        name=resume.name,
        email=resume.email,
        phone=resume.phone,
        score=score,
        verdict=get_verdict(score),
        matched_skills=match.matched_skills,
        missing_skills=match.missing_skills,
        experience_match=match.experience_match,
        education_match=match.education_match,
        reasons=match.reasons,
    )


@app.get("/api/health")
def health():
    return {"status": "ok", "model": MODEL_NAME}


@app.post("/api/screen", response_model=ScreenResponse)
async def screen(
    job_description: str = Form(...),
    files: list[UploadFile] = File(...),
    user_id: str = Form("default"),
):
    if not job_description.strip():
        raise HTTPException(status_code=400, detail="Job description is required.")

    if not files:
        raise HTTPException(status_code=400, detail="Upload at least one resume.")

    if len(files) > MAX_RESUMES:
        raise HTTPException(
            status_code=400,
            detail=f"Maximum {MAX_RESUMES} resumes allowed per request.",
        )

    check_rate_limit(user_id)

    client = get_client()
    job = parse_job_description(client, job_description)

    results = []
    skipped = []

    for upload in files:
        data = await upload.read()

        if len(data) / (1024 * 1024) > MAX_FILE_SIZE_MB:
            skipped.append(f"{upload.filename}: exceeds the {MAX_FILE_SIZE_MB} MB limit.")
            continue

        try:
            resume_text = read_resume(upload.filename, data)
        except HTTPException as error:
            skipped.append(str(error.detail))
            continue
        except Exception:
            skipped.append(f"{upload.filename}: could not be read.")
            continue

        if not resume_text.strip():
            skipped.append(f"{upload.filename}: no readable text found.")
            continue

        try:
            resume = parse_resume(client, resume_text)
            match = match_resume(client, job, resume)
        except Exception:
            skipped.append(f"{upload.filename}: screening failed, try again.")
            continue

        results.append(create_final_result(upload.filename, resume, match))

    results.sort(key=lambda result: result.score, reverse=True)

    return ScreenResponse(
        role=job.role,
        required_skills=job.required_skills,
        results=results,
        skipped=skipped,
    )
BROWSER_HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/122.0 Safari/537.36",
    "Accept-Language": "en-US,en;q=0.9",
}

FETCH_FAILED = (
    "That link could not be opened. Many job boards block automated access "
    "or need a login. Paste the description text instead."
)

THIN_PAGE = (
    "The page opened but did not contain enough readable text. "
    "Paste the description text instead."
)
NOISE_LINES = {
    "skip to main content",
    "skip to footer",
    "careers",
    "expand menu",
    "save this job",
    "unsave this job",
    "share this job",
    "copy link",
    "linkedin",
    "x",
    "facebook",
    "email",
    "apply for this job",
}

NOISE_PREFIXES = (
    "learn more",
    "discover where this job fits",
    "about accenture",
    "additional information",
    "equal employment opportunity",
    "job candidates will not be",
    "accenture is committed",
    "please read accenture",
    "we work with one shared purpose",
    "we believe that delivering value",
    "at accenture, we see well-being",
    "join accenture to work",
)


def extract_page_text(html):
    soup = BeautifulSoup(html, "html.parser")

    for tag in soup(["script", "style", "noscript", "header", "footer", "nav", "svg", "button", "a"]):
        tag.decompose()

    title = soup.title.get_text(strip=True) if soup.title else None

    raw_lines = soup.get_text("\n").split("\n")
    cleaned = []
    seen = set()

    for line in raw_lines:
        stripped = line.strip()

        if not stripped:
            continue

        lowered = stripped.lower()

        if lowered in NOISE_LINES:
            continue

        if lowered.startswith(NOISE_PREFIXES):
            continue

        if len(stripped) < 3:
            continue

        if stripped in seen:
            continue

        seen.add(stripped)
        cleaned.append(stripped)

    text = "\n".join(cleaned)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)

    return title, text.strip()

@app.post("/api/fetch-job", response_model=JobUrlResponse)
def fetch_job(payload: JobUrlRequest):
    url = payload.url.strip()

    if not url:
        raise HTTPException(status_code=400, detail="Enter a link first.")

    if not url.startswith(("http://", "https://")):
        url = "https://" + url

    try:
        with httpx.Client(
            timeout=15.0, follow_redirects=True, headers=BROWSER_HEADERS
        ) as http_client:
            response = http_client.get(url)
            response.raise_for_status()
    except Exception:
        raise HTTPException(status_code=422, detail=FETCH_FAILED)

    if "html" not in response.headers.get("content-type", "") and not response.text:
        raise HTTPException(status_code=422, detail=FETCH_FAILED)

    try:
        title, text = extract_page_text(response.text)
    except Exception:
        raise HTTPException(status_code=422, detail=FETCH_FAILED)

    if len(text) < 250:
        raise HTTPException(status_code=422, detail=THIN_PAGE)

    return JobUrlResponse(url=url, title=title, job_description=text[:20000])